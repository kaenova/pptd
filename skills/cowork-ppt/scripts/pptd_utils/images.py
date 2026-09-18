"""Image element: fit (cover/contain/fill), crop (incl. negative outset),
cropShape clipping, border, shadow, opacity, rotation/flip."""
import hashlib
from pathlib import Path

from PIL import Image as PILImage
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Pt

from .shapes import apply_preset_geometry, build_custgeom, _apply_alpha
from .style import fetch_url, cover_src_rect, ln_xml, shadow_xml
from .xmlutil import sub, set_xfrm_flip_rot, insert_ordered, SPPR_ORDER

PREP_CACHE = Path.home() / ".cache" / "pptd_utils" / "prepped"


def _prep_negative_crop(src, crop):
    """Negative crop values -> expand canvas with transparency; returns path."""
    l = crop.get("left", 0) or 0
    t = crop.get("top", 0) or 0
    r = crop.get("right", 0) or 0
    b = crop.get("bottom", 0) or 0
    if min(l, t, r, b) >= 0:
        return src
    key = hashlib.sha1(f"{src}|{l},{t},{r},{b}".encode()).hexdigest()[:16]
    dst = PREP_CACHE / f"{key}.png"
    if not dst.exists():
        PREP_CACHE.mkdir(parents=True, exist_ok=True)
        with PILImage.open(src) as im:
            im = im.convert("RGBA")
            w, h = im.size
            pad = (max(0, -l) * w, max(0, -t) * h, max(0, -r) * w, max(0, -b) * h)
            canvas = PILImage.new("RGBA", (w + int(pad[0] + pad[2]), h + int(pad[1] + pad[3])), (0, 0, 0, 0))
            canvas.paste(im, (int(pad[0]), int(pad[1])))
            canvas.save(dst)
    return dst


def render_image(slide, el, theme, root):
    src = el["src"]
    path = fetch_url(src) if str(src).startswith(("http://", "https://")) else (root / src)
    x, y, w, h = el["bounds"]
    crop = el.get("crop") or {}
    if isinstance(crop, (list, tuple)):
        crop = dict(zip(("left", "top", "right", "bottom"), crop))
    path = _prep_negative_crop(path, crop)
    pos_crop = {k: max(0.0, v) for k, v in crop.items() if (v or 0) > 0} if crop else {}

    fite = el.get("fit") or {}
    fit = fite.get("mode", "cover") if isinstance(fite, dict) else str(fite)
    px, py, pw, ph = x, y, w, h
    src_rect = None
    if fit == "cover":
        src_rect = cover_src_rect(path, w, h, pos_crop or None)
    elif fit == "contain":
        try:
            with PILImage.open(path) as im:
                iw, ih = im.size
        except Exception:
            iw = ih = 1
        if iw and ih and w and h:
            if iw / ih > w / h:
                pw = w
                ph = w * ih / iw
            else:
                ph = h
                pw = h * iw / ih
            px = x + (w - pw) / 2
            py = y + (h - ph) / 2

    pic = slide.shapes.add_picture(str(path), Pt(px), Pt(py), Pt(pw), Pt(ph))
    blipFill = pic._element.blipFill
    if src_rect:
        st = blipFill.find(qn("a:stretch"))
        sr = sub(blipFill, "a:srcRect", src_rect)
        if st is not None:
            st.addprevious(sr)
    opacity = el.get("opacity", 1.0)
    if opacity is not None and float(opacity) < 1.0:
        blip = blipFill.find(qn("a:blip"))
        sub(blip, "a:alphaModFix", {"amt": str(int(float(opacity) * 100000))})

    spPr = pic._element.spPr
    cs = el.get("cropShape")
    if isinstance(cs, str):
        cs = {"shapeName": cs}
    if cs:
        name = cs.get("shapeName")
        if name == "custom":
            build_custgeom(spPr, cs.get("viewBox", [pw, ph]), cs.get("path", ""))
        else:
            geom = spPr.find(qn("a:prstGeom"))
            if geom is None:
                geom = insert_ordered(spPr, sub(spPr, "a:prstGeom", {"prst": "rect"}), SPPR_ORDER)
            else:
                geom.set("prst", name or "rect")
                avLst = geom.find(qn("a:avLst"))
                for gd in list(avLst):
                    avLst.remove(gd)
            apply_preset_geometry(pic, name, cs.get("adjustments"))
    ln_xml(spPr, el.get("border"), theme)
    shadow_xml(spPr, el.get("shadow"), theme)
    set_xfrm_flip_rot(pic, el.get("rotation") or 0, el.get("flip"))
    return pic

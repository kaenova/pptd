"""Text element + shared run emission (used by text boxes and table cells)."""
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from .rich import parse_rich, LIST_TYPES
from .style import resolve_color, color_el, css_px
from .xmlutil import sub, insert_ordered, RPR_ORDER, set_xfrm_flip_rot

H_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
           "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY,
           "distributed": PP_ALIGN.DISTRIBUTE}
V_ALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
           "bottom": MSO_ANCHOR.BOTTOM}

TEXT_DEFAULTS = {"color": "#000000", "fontSize": 18, "fontFamily": "MiSans",
                 "bold": False, "italic": False, "lineHeight": 1,
                 "letterSpacing": 0, "marginTop": 0}

TEXTSTYLE_KEYS = ("color", "fontSize", "fontFamily", "bold", "italic",
                  "backgroundColor", "lineHeight", "lineHeightPx",
                  "letterSpacing", "marginTop")


def resolve_base_style(c, theme, extra=None):
    """Style priority chain for TextContent-like dicts -> flat style dict."""
    base = dict(TEXT_DEFAULTS)
    if c.get("style"):
        ref = theme.get("textStyles", {}).get(str(c["style"]).lstrip("$"), {})
        for k in TEXTSTYLE_KEYS:
            if ref.get(k) is not None:
                base[k] = ref[k]
    for k in TEXTSTYLE_KEYS:
        if c.get(k) is not None:
            base[k] = c[k]
    for k, v in (extra or {}).items():
        if v is not None:
            base[k] = v
    return base


def _font_family_names(ff):
    """FontFamily -> (latin, ea)."""
    if isinstance(ff, dict):
        return ff.get("latin"), ff.get("ea")
    return ff, ff


def emit_runs(p, para, base, theme, opacity=1.0):
    """Emit parsed runs of one paragraph into a pptx paragraph p."""
    for run in para.runs:
        if run.text == "\n":
            p.add_line_break()
            continue
        if run.math is not None:
            omml = latex_to_omml(run.math)
            if omml is not None:
                p._p.append(omml)
                continue
        if not run.text:
            continue
        r = p.add_run()
        r.text = run.text
        span = run.span
        f = r.font
        size = css_px(span.get("font-size")) or base.get("fontSize", 18)
        f.size = Pt(size)
        latin, ea = _font_family_names(
            span.get("font-family") or base.get("fontFamily", "MiSans"))
        f.name = str(latin or "MiSans")
        rPr = r._r.get_or_add_rPr()
        if ea:
            ea_el = rPr.find(qn("a:ea"))
            if ea_el is None:
                ea_el = insert_ordered(rPr, etree.Element(qn("a:ea")), RPR_ORDER)
            ea_el.set("typeface", str(ea))
        f.bold = bool(base.get("bold")) or run.bold
        f.italic = bool(base.get("italic")) or run.italic
        if run.underline:
            f.underline = True
        if run.strike:
            rPr.set("strike", "sngStrike")
        if run.sup:
            rPr.set("baseline", "30000")
        if run.sub:
            rPr.set("baseline", "-25000")
        ls = base.get("letterSpacing")
        if ls:
            rPr.set("spc", str(int(float(ls) * 100)))
        # color (+alpha)
        color = span.get("color") or base.get("color", "#000000")
        if run.link and not span.get("color"):
            color = "#0563C1"
            f.underline = True
        rgb, alpha = resolve_color(color, theme, opacity)
        f.color.rgb = rgb
        if alpha < 1.0:
            srgb = rPr.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
            if srgb is None:
                srgb = rPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            if srgb is not None:
                sub(srgb, "a:alpha", {"val": str(int(alpha * 100000))})
        if run.link:
            r.hyperlink.address = run.link
        bg = span.get("background-color")
        if bg:
            hl = etree.Element(qn("a:highlight"))
            color_el(hl, bg, theme)
            insert_ordered(rPr, hl, RPR_ORDER)


def _apply_para_style(p, para, base, theme, def_align):
    pPr = p._p.get_or_add_pPr()
    p.alignment = H_ALIGN.get(para.align or def_align, PP_ALIGN.LEFT)
    lh_px = para.lh_px if para.lh_px is not None else (
        css_px(base.get("lineHeightPx")) if base.get("lineHeightPx") else None)
    if lh_px:
        p.line_spacing = Pt(lh_px)
    else:
        lh = para.lh if para.lh is not None else base.get("lineHeight", 1)
        if lh and float(lh) != 1.0:
            p.line_spacing = float(lh)
    mt = para.mt if para.mt is not None else base.get("marginTop", 0)
    if mt:
        p.space_before = Pt(float(mt))
    ml = para.ml
    if ml:
        pPr.set("marL", str(int(Pt(float(ml)))))
    if para.mr:
        pPr.set("marR", str(int(Pt(float(para.mr)))))
    if para.ls:
        # li-level letter-spacing: applied via runs below (base override)
        base["letterSpacing"] = para.ls
    # bullets
    if para.bullet or para.ordered:
        lsty = para.list_style or "disc"
        if para.ordered and lsty == "disc":
            lsty = "decimal"
        if lsty in (None, "none"):
            insert_ordered(pPr, etree.SubElement(pPr, qn("a:buNone")), _PPR_BUL)
        elif lsty in LIST_TYPES and isinstance(LIST_TYPES[lsty], str) and LIST_TYPES[lsty][0].isalpha():
            bul = etree.SubElement(pPr, qn("a:buAutoNum"), {"type": LIST_TYPES[lsty]})
            insert_ordered(pPr, bul, _PPR_BUL)
            _add_bu_font(pPr)
        else:
            char = LIST_TYPES.get(lsty, "\u2022") if lsty in LIST_TYPES else lsty
            bul = etree.SubElement(pPr, qn("a:buChar"), {"char": str(char)})
            insert_ordered(pPr, bul, _PPR_BUL)
            _add_bu_font(pPr)
        if pPr.get("marL") is None:
            pPr.set("marL", "228600")
        pPr.set("indent", "-228600")


_PPR_BUL = [qn(t) for t in ("a:buClrTx", "a:buClr", "a:buSzTx", "a:buSzPct",
                            "a:buSzPts", "a:buFontTx", "a:buFont", "a:buNone",
                            "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst",
                            "a:defRPr", "a:extLst")]


def _add_bu_font(pPr):
    bf = etree.SubElement(pPr, qn("a:buFont"), {"typeface": "Arial"})
    insert_ordered(pPr, bf, _PPR_BUL)


def emit_text(tf, content, theme, default_align=("left", "top"), opacity=1.0,
              extra_style=None, vertical=None, wrap=None):
    """Emit TextContent into a pptx text_frame tf."""
    base = resolve_base_style(content, theme, extra_style)
    align = content.get("align") or list(default_align)
    tf.word_wrap = content.get("wrap", True) if wrap is None else wrap
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = V_ALIGN.get(align[1], MSO_ANCHOR.TOP)
    if vertical or content.get("textDirection") == "vertical":
        tf._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
    paras = parse_rich(content.get("text", ""))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _apply_para_style(p, para, base, theme, align[0])
        emit_runs(p, para, base, theme, opacity)
    return base


def render_text(slide, el, theme):
    c = el["content"]
    x, y, w, h = el["bounds"]
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.auto_size = None
    base = emit_text(tf, c, theme, opacity=el.get("opacity", 1.0))
    rotation = el.get("rotation") or 0
    flip = el.get("flip")
    if rotation or flip:
        set_xfrm_flip_rot(tb, rotation, flip)
    # text gradient
    grad = c.get("gradient")
    if grad:
        _text_gradient(tf, grad, theme)
    shadow = c.get("shadow")
    if shadow:
        _text_shadow(tf, shadow, theme)
    return tb


def _text_gradient(tf, grad, theme):
    from .xmlutil import RPR_ORDER as ORDER
    from .style import fill_xml
    for para in tf.paragraphs:
        for run in para.runs:
            rPr = run._r.get_or_add_rPr()
            fill_xml(rPr, grad, theme, order=ORDER)


def _text_shadow(tf, shadow, theme):
    from .style import shadow_xml
    for para in tf.paragraphs:
        for run in para.runs:
            rPr = run._r.get_or_add_rPr()
            shadow_xml(rPr, shadow, theme, order=RPR_ORDER)


def latex_to_omml(latex):
    """Optional: LaTeX -> m:oMath element via latex2mathml + mathml2omml."""
    try:
        import latex2mathml.converter as l2m
        import mathml2omml
    except ImportError:
        return None
    try:
        mathml = l2m.convert(latex)
        omml = mathml2omml.convert(mathml)
        if isinstance(omml, bytes):
            omml = omml.decode()
        ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
        root = etree.fromstring("<root xmlns:m='%s'>%s</root>" % (ns["m"], omml))
        return root.find("m:oMath", ns)
    except Exception:
        return None

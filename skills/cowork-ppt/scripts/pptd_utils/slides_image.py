"""Slide image export: collage with numbering + single-slide PNGs.

Renders via the built .pptx: each wanted slide is isolated into a
one-slide package and exported by LibreOffice headless (soffice).
Pillow composes the collage with numbered badges.

ponytail: one soffice run per slide (~1.5s each); switch to a single
PDF export + PyMuPDF rasterize when decks get big.
"""
import math
import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from .main import build
from .charts_raster import _font  # system TTF lookup, shared


def _soffice_png(pptx: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "png",
         str(pptx), "--outdir", str(out_dir)],
        check=True, capture_output=True, timeout=120)
    png = out_dir / (pptx.stem + ".png")
    if not png.exists():
        raise RuntimeError(f"soffice did not produce {png}")
    return png


def _split(pptx: Path, index: int, dst: Path) -> Path:
    """index 1-based -> one-slide pptx at dst."""
    from pptx import Presentation
    prs = Presentation(str(pptx))
    lst = prs.slides._sldIdLst
    keep = list(lst)[index - 1]
    for s in list(lst):
        if s is not keep:
            lst.remove(s)
    prs.save(str(dst))
    return dst


def render_slide_images(pptd_path, out_dir, slides=None, tmp_pptx=None):
    """-> [(slide_no, Path)] for the requested 1-based slide numbers."""
    pptd_path = Path(pptd_path)
    out_dir = Path(out_dir)
    if tmp_pptx is None:
        tmp_pptx = Path(tempfile.mkdtemp(prefix="pptd_utils-img-")) / "deck.pptx"
        build(pptd_path, tmp_pptx)
    if slides is None:
        from pptx import Presentation
        slides = range(1, len(Presentation(str(tmp_pptx)).slides._sldIdLst) + 1)
    work = Path(tempfile.mkdtemp(prefix="pptd_utils-split-"))
    results = []
    for n in slides:
        one = _split(tmp_pptx, n, work / f"s{n}.pptx")
        png = _soffice_png(one, work)
        final = out_dir / f"slide_{n:02d}.png"
        out_dir.mkdir(parents=True, exist_ok=True)
        shutil.move(str(png), final)
        results.append((n, final))
    shutil.rmtree(work, ignore_errors=True)
    return results


def render_slide(pptd_path, slide, out_png):
    """Export one slide (1-based) as a PNG."""
    n, png = render_slide_images(pptd_path, Path(out_png).parent, [slide])[0]
    dst = Path(out_png)
    if png != dst:
        shutil.move(str(png), dst)
    return dst


def render_collage(pptd_path, out_png, tile_w=480, gap=20, numbering=True,
                   bg="#1F2430", label="Slide", slides=None, tmp_pptx=None):
    """Grid collage of all (or selected) slides with numbered badges."""
    import tempfile
    tmp = Path(tempfile.mkdtemp(prefix="pptd_utils-col-"))
    try:
        imgs = render_slide_images(pptd_path, tmp, slides, tmp_pptx=tmp_pptx)
        tiles = [(n, Image.open(p).convert("RGB")) for n, p in imgs]
        if not tiles:
            raise ValueError("no slides rendered")
        n = len(tiles)
        cols = math.ceil(math.sqrt(n))
        rows = math.ceil(n / cols)
        tw = tile_w
        th = round(tiles[0][1].height * tw / tiles[0][1].width)
        badge = max(18, tw // 18)
        footer = badge + 14 if numbering else 0
        W = cols * tw + (cols + 1) * gap
        H = rows * (th + footer) + (rows + 1) * gap
        sheet = Image.new("RGB", (W, H), _rgb(bg))
        d = ImageDraw.Draw(sheet)
        f = _font(round(badge * 0.62))
        for i, (no, im) in enumerate(tiles):
            r, c = divmod(i, cols)
            x = gap + c * (tw + gap)
            y = gap + r * (th + footer + gap)
            sheet.paste(im.resize((tw, th), Image.LANCZOS), (x, y))
            d.rectangle([x, y, x + tw, y + th], outline=(255, 255, 255),
                        width=1)
            if numbering:
                text = f"{label} {no}"
                bb = d.textbbox((0, 0), text, font=f)
                tw_t, th_t = bb[2] - bb[0], bb[3] - bb[1]
                cx = x + (tw - tw_t) / 2 - bb[0]
                cy = y + th + (footer - th_t) / 2 - bb[1]
                d.ellipse([cx - badge * 0.9, cy - badge * 0.45,
                           cx + tw_t + badge * 0.9, cy + th_t + badge * 0.45],
                          fill=(37, 99, 235))
                d.text((cx, cy), text, font=f, fill=(255, 255, 255))
        out = Path(out_png)
        out.parent.mkdir(parents=True, exist_ok=True)
        sheet.save(out)
        return out
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _rgb(hexs):
    h = hexs.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))

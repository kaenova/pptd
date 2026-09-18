"""Line element: bezier/rounded/sharp polylines via custGeom + arrowheads."""
import math

from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.util import Pt
from lxml import etree

from .shapes import build_custgeom, _ipf
from .style import ln_xml, shadow_xml
from .xmlutil import sub, set_xfrm_flip_rot

ARROW_MAP = {"arrow": "triangle", "stealth": "stealth",
             "diamond": "diamond", "oval": "oval"}


def _pts(points):
    toks = str(points).replace(",", " ").split()
    out = []
    for i in range(0, len(toks) - 1, 2):
        out.append((float(toks[i]), float(toks[i + 1])))
    return out


def _path_cmds(pts, curve):
    """Build (tag, vals) list in viewBox coords from point list."""
    n = len(pts)
    cmds = [("M", [pts[0][0], pts[0][1]])]
    if curve is None or curve == "sharp" or n == 2:
        for p in pts[1:]:
            cmds.append(("L", [p[0], p[1]]))
        return cmds
    if curve == "smooth":
        i = 1
        while i < n:
            if n - i >= 3:
                cmds.append(("C", [pts[i][0], pts[i][1], pts[i + 1][0],
                                   pts[i + 1][1], pts[i + 2][0], pts[i + 2][1]]))
                i += 3
            elif n - i == 2:
                # quadratic through one control
                x0, y0 = pts[i - 1]
                x1, y1 = pts[i]
                x2, y2 = pts[i + 1]
                cx1 = x0 + 2 / 3 * (x1 - x0)
                cy1 = y0 + 2 / 3 * (y1 - y0)
                cx2 = x2 + 2 / 3 * (x1 - x2)
                cy2 = y2 + 2 / 3 * (y1 - y2)
                cmds.append(("C", [cx1, cy1, cx2, cy2, x2, y2]))
                i += 2
            else:
                cmds.append(("L", [pts[i][0], pts[i][1]]))
                i += 1
        return cmds
    # round: polyline with corner fillets
    r = None
    for i in range(1, n - 1):
        e1 = math.hypot(pts[i][0] - pts[i - 1][0], pts[i][1] - pts[i - 1][1])
        e2 = math.hypot(pts[i + 1][0] - pts[i][0], pts[i + 1][1] - pts[i][1])
        cand = min(e1, e2) / 3.0
        r = cand if r is None else min(r, cand)
    r = min(r or 0, 12.0)
    if r < 0.01:
        for p in pts[1:]:
            cmds.append(("L", [p[0], p[1]]))
        return cmds
    cur = pts[0]
    for i in range(1, n - 1):
        p, nx = pts[i], pts[i + 1]
        d1 = (p[0] - cur[0], p[1] - cur[1])
        l1 = math.hypot(*d1) or 1.0
        d2 = (nx[0] - p[0], nx[1] - p[1])
        l2 = math.hypot(*d2) or 1.0
        r1 = min(r, l1 / 2)
        r2 = min(r, l2 / 2)
        c1 = (p[0] - d1[0] / l1 * r1, p[1] - d1[1] / l1 * r1)
        c2 = (p[0] + d2[0] / l2 * r2, p[1] + d2[1] / l2 * r2)
        cmds.append(("L", [c1[0], c1[1]]))
        cmds.append(("Q", [p[0], p[1], c2[0], c2[1]]))
        cur = c2
    last = pts[-1]
    cmds.append(("L", [last[0], last[1]]))
    return cmds


def _emit_path(path_el, cmds):
    for tag, vals in cmds:
        if tag == "M":
            mt = sub(path_el, "a:moveTo")
            pt = sub(mt, "a:pt", {"x": _i(vals[0]), "y": _i(vals[1])})
        elif tag == "L":
            lt = sub(path_el, "a:lnTo")
            sub(lt, "a:pt", {"x": _i(vals[0]), "y": _i(vals[1])})
        elif tag == "C":
            cb = sub(path_el, "a:cubicBezTo")
            for j in (0, 2, 4):
                sub(cb, "a:pt", {"x": _i(vals[j]), "y": _i(vals[j + 1])})
        elif tag == "Q":
            qb = sub(path_el, "a:quadBezTo")
            for j in (0, 2):
                sub(qb, "a:pt", {"x": _i(vals[j]), "y": _i(vals[j + 1])})


def _i(v):
    return str(int(round(float(v))))


def render_line(slide, el, theme):
    pts = _pts(el.get("points", ""))
    if len(pts) < 2:
        raise ValueError("line needs >= 2 points")
    x, y, w, h = el["bounds"]
    vw, vh = el.get("viewBox", [w, h])

    # Use a native connector for ordinary 2-point lines. PowerPoint can drop
    # custom-geometry line shapes even though cowork/browser renderers display them.
    if len(pts) == 2:
        sx, sy = pts[0]
        ex, ey = pts[1]
        x1, y1 = x + sx / vw * w, y + sy / vh * h
        x2, y2 = x + ex / vw * w, y + ey / vh * h
        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        shape.shadow.inherit = False
        spPr = shape._element.spPr
        border = dict(el.get("border") or {})
        border.setdefault("width", 2)
        border.setdefault("color", "#000000")
        ln = ln_xml(spPr, border, theme)
        arrows = el.get("arrow") or [None, None]
        if arrows[0]:
            sub(ln, "a:headEnd", {"type": ARROW_MAP.get(arrows[0], "triangle"),
                                  "w": "med", "len": "med"})
        if arrows[1]:
            sub(ln, "a:tailEnd", {"type": ARROW_MAP.get(arrows[1], "triangle"),
                                  "w": "med", "len": "med"})
        set_xfrm_flip_rot(shape, el.get("rotation") or 0, el.get("flip"))
        return shape

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.shadow.inherit = False
    spPr = shape._element.spPr
    # geometry
    for t in ("a:prstGeom", "a:custGeom"):
        e = spPr.find(qn(t))
        if e is not None:
            spPr.remove(e)
    cust = etree.Element(qn("a:custGeom"))
    spPr.insert(0, cust)
    sub(cust, "a:avLst")
    sub(cust, "a:gdLst")
    pathLst = sub(cust, "a:pathLst")
    path = sub(pathLst, "a:path", {"w": _i(vw), "h": _i(vh)})
    _emit_path(path, _path_cmds(pts, el.get("curve", "round")))
    # no fill, line style
    for e in spPr.findall(qn("a:noFill")):
        spPr.remove(e)
    from .xmlutil import SPPR_ORDER, insert_ordered
    insert_ordered(spPr, etree.Element(qn("a:noFill")), SPPR_ORDER)
    border = el.get("border") or {}
    border = dict(border)
    border.setdefault("width", 2)
    border.setdefault("color", "#000000")
    ln = ln_xml(spPr, border, theme)
    arrows = el.get("arrow") or [None, None]
    if arrows[0]:
        sub(ln, "a:headEnd", {"type": ARROW_MAP.get(arrows[0], "triangle"),
                              "w": "med", "len": "med"})
    if arrows[1]:
        sub(ln, "a:tailEnd", {"type": ARROW_MAP.get(arrows[1], "triangle"),
                               "w": "med", "len": "med"})
    opacity = el.get("opacity", 1.0)
    if opacity is not None and float(opacity) < 1.0:
        val = str(int(float(opacity) * 100000))
        for srgb in spPr.iter(qn("a:srgbClr")):
            if srgb.find(qn("a:alpha")) is None:
                sub(srgb, "a:alpha", {"val": val})
    shadow_xml(spPr, el.get("shadow"), theme)
    set_xfrm_flip_rot(shape, el.get("rotation") or 0, el.get("flip"))
    return shape
